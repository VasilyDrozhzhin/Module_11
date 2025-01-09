
import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import load_workbook

# URL страницы с турнирной таблицей
url = 'https://www.sports.ru/basketball/tournament/nba/table/'

response = requests.get(url)

if response.status_code == 200:
    soup = BeautifulSoup(response.text, 'html.parser')

    tables = soup.find_all('table')

    if len(tables) >= 3:
        with pd.ExcelWriter('nba_tables.xlsx', engine='xlsxwriter') as writer:
            second_table = tables[1]
            if second_table:
                headers = [header.text.strip() for header in second_table.find_all('th')]
                if not headers:
                    num_columns = len(second_table.find_all('tr')[0].find_all('td'))
                    headers = [f"Column {i + 1}" for i in range(num_columns)]
                rows = []
                for row in second_table.find_all('tr'):
                    columns = row.find_all('td')
                    if columns:
                        data = [column.text.strip() for column in columns]
                        rows.append(data)
                df_second = pd.DataFrame(rows, columns=headers)
                df_second.to_excel(writer, sheet_name='Восточная конференция', index=False)

            third_table = tables[2]
            if third_table:
                headers = [header.text.strip() for header in third_table.find_all('th')]
                if not headers:
                    num_columns = len(third_table.find_all('tr')[0].find_all('td'))
                    headers = [f"Column {i + 1}" for i in range(num_columns)]
                rows = []
                for row in third_table.find_all('tr'):
                    columns = row.find_all('td')
                    if columns:
                        data = [column.text.strip() for column in columns]
                        rows.append(data)
                df_third = pd.DataFrame(rows, columns=headers)
                df_third.to_excel(writer, sheet_name='Западная конференция', index=False)

        # Открываем созданный файл Excel для редактирования
        workbook = load_workbook('nba_tables.xlsx')

        # Удаляем первую строку на каждом листе
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet.delete_rows(1)

        workbook.save('nba_tables.xlsx')

        print("Данные второй и третьей таблиц успешно экспортированы в nba_tables.xlsx")
    else:
        print("На странице меньше трех таблиц.")
else:
    print(f"Ошибка при загрузке страницы: {response.status_code}")
